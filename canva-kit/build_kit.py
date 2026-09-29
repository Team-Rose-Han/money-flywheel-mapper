"""Builds the Canva-ready Money Flywheel Map kit (money-flywheel-map.pptx).

Every element is a native shape: cards are rounded rectangles with their label
typed inside, connectors are plain lines, dots are circles. Canva imports a
.pptx as separate editable elements, so text, card size and colour can all be
changed there. Run: python3 build_kit.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PX = 9525  # EMU per pixel at 96 dpi, so coordinates below are in pixels

NAVY = RGBColor(0x0B, 0x1D, 0x80)
LINE = RGBColor(0x1E, 0x3A, 0xB8)
INK = RGBColor(0x14, 0x14, 0x14)
LIME = RGBColor(0xD9, 0xFF, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Poppins"

LINE_W = 3
CARD_BORDER = 3
RADIUS = 22


def px(v):
    return Emu(int(v * PX))


def style_text(frame, text, size_pt, color, spacing_pt, bold=True,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = px(10)
    frame.margin_top = frame.margin_bottom = px(4)
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = color
    run._r.get_or_add_rPr().set("spc", str(int(spacing_pt * 100)))


def card(slide, name, x, y, w, h, text, fill=WHITE, text_color=INK,
         border=True, radius=RADIUS, size=18, spacing=4.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    s.name = name
    s.adjustments[0] = radius / min(w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if border:
        s.line.color.rgb = INK
        s.line.width = px(CARD_BORDER)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    style_text(s.text_frame, text, size, text_color, spacing)
    return s


def line(slide, name, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.name = name
    c.line.color.rgb = LINE
    c.line.width = px(LINE_W)
    return c


def dot(slide, name, cx, cy, d=14):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - d / 2), px(cy - d / 2), px(d), px(d))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = LINE
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def arrowhead(slide, name, tip_x, tip_y, size=16):
    # Open chevron made of two lines, so it survives import into any tool.
    line(slide, name + " top", tip_x - size * 0.8, tip_y - size * 0.75, tip_x, tip_y)
    line(slide, name + " bottom", tip_x - size * 0.8, tip_y + size * 0.75, tip_x, tip_y)


def text_box(slide, name, x, y, w, h, text, size, color, spacing,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    t = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    t.name = name
    style_text(t.text_frame, text, size, color, spacing, align=align, anchor=anchor)
    return t


def background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def build_map(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)

    text_box(s, "Title", 14, 30, 1300, 70, title, 34, INK, 2)
    card(s, "Pill: Your map", 1810, 33, 160, 51, "YOUR MAP", fill=LIME,
         border=False, radius=25.5, size=13, spacing=2.5)

    # Business side
    card(s, "Card: Revenue", 66, 150, 271, 102, "REVENUE", fill=NAVY,
         text_color=WHITE, border=False, size=17, spacing=4.5)
    card(s, "Card: Biz Hub", 32, 357, 339, 148, "BIZ HUB")
    card(s, "Card: Taxes", 32, 611, 339, 127, "TAXES")
    line(s, "Line: Revenue to Biz Hub", 201.5, 252, 201.5, 357)
    dot(s, "Dot: Revenue to Biz Hub", 201.5, 304)
    line(s, "Line: Biz Hub to Taxes", 201.5, 505, 201.5, 611)
    dot(s, "Dot: Biz Hub to Taxes", 201.5, 558)

    # Pay yourself
    line(s, "Line: Pay yourself", 371, 431, 1062, 431)
    arrowhead(s, "Arrow: Pay yourself", 1062, 431)
    text_box(s, "Label: Pay yourself", 516, 388, 400, 36, "PAY YOURSELF", 12.5,
             LINE, 3.5, align=PP_ALIGN.CENTER)

    # Personal side
    card(s, "Card: Personal Hub", 1064, 357, 340, 148, "PERSONAL HUB")
    line(s, "Line: Personal Hub down", 1234, 505, 1234, 558)
    line(s, "Line: Bucket rail", 676, 558, 1792, 558)
    dot(s, "Dot: Personal Hub split", 1234, 558)

    buckets = [("Bills", 506), ("Spending", 878), ("Upcoming", 1250),
               ("Financial Goals", 1622)]
    for label, x in buckets:
        cx = x + 170
        line(s, f"Line: to {label}", cx, 558, cx, 611)
        card(s, f"Card: {label}", x, 611, 340, 268, label.upper())
    return s


def build_parts(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    text_box(s, "Title", 14, 30, 1300, 70, "SPARE PARTS", 34, INK, 2)
    text_box(s, "Note", 24, 100, 1500, 40,
             "Copy these onto the map to add accounts. Hold Shift while dragging a line end to keep it straight.",
             12, INK, 0.5, anchor=MSO_ANCHOR.TOP)

    card(s, "Card: Hub size", 40, 180, 340, 148, "NEW HUB")
    card(s, "Card: Bucket size", 440, 180, 340, 268, "NEW BUCKET")
    card(s, "Card: Small size", 840, 180, 339, 127, "NEW ACCOUNT")
    card(s, "Card: Navy", 1240, 180, 271, 102, "INCOME", fill=NAVY,
         text_color=WHITE, border=False, size=17)
    card(s, "Pill: Lime", 1580, 180, 160, 51, "YOUR MAP", fill=LIME,
         border=False, radius=25.5, size=13, spacing=2.5)

    line(s, "Line: horizontal", 40, 560, 500, 560)
    arrowhead(s, "Arrow: horizontal", 500, 560)
    text_box(s, "Label: flow", 120, 517, 300, 36, "FLOW LABEL", 12.5, LINE, 3.5,
             align=PP_ALIGN.CENTER)
    line(s, "Line: vertical", 620, 500, 620, 640)
    dot(s, "Dot", 620, 570)
    line(s, "Line: plain vertical", 740, 500, 740, 640)
    dot(s, "Dot: spare", 860, 570)


def main():
    prs = Presentation()
    prs.slide_width = px(2000)
    prs.slide_height = px(940)
    build_map(prs, "ROSE'S MONEY FLYWHEEL MAP")
    build_map(prs, "YOUR MONEY FLYWHEEL MAP")
    build_parts(prs)
    prs.save("money-flywheel-map.pptx")


if __name__ == "__main__":
    main()
